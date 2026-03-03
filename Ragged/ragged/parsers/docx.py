from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from ragged.models import Document


class DocxParser:
    """Parse DOCX files, extracting paragraph text."""

    def parse(self, file_path: Path) -> list[Document]:
        file_path = Path(file_path)
        doc = DocxDocument(str(file_path))

        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)

        text = "\n\n".join(paragraphs)

        # Try to get title from core properties
        title = doc.core_properties.title if doc.core_properties.title else None

        metadata = {
            "source": str(file_path),
            "page": None,
            "section": title,
        }

        return [Document(text=text, metadata=metadata)] if text.strip() else []
